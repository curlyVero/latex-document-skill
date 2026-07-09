from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

out=Path('output'); out.mkdir(exist_ok=True)

def fmt(r,size=12,bold=False):
    r.font.name='宋体'
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    r.font.size=Pt(size); r.bold=bold; r.font.highlight_color=WD_COLOR_INDEX.YELLOW

def add(doc,t,size=12,b=False):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.3
    r=p.add_run(t); fmt(r,size,b); return p

def head(doc,t):
    r=doc.add_paragraph().add_run(t); fmt(r,14,True)

def make_report():
    d=Document()
    d.add_paragraph('铁路单北斗差分定位算法优化与应用研究').alignment=WD_ALIGN_PARAGRAPH.CENTER
    add(d,'企业科研结题报告优化修订稿（标黄修改版）',16,True)
    add(d,'说明：黄色内容为针对原报告进行的逻辑、表达、成果总结和验收导向优化内容。原实验数据、公式和图表建议保留。',11)
    head(d,'一、摘要优化')
    add(d,'本课题面向铁路行业单北斗规模应用需求，围绕自主可控高精度定位关键技术开展研究。针对铁路地基增强系统对多系统GNSS依赖、复杂环境下单北斗定位可靠性不足等问题，开展单北斗时空基准评估、服务端改造、电离层精化建模、闪烁监测、接收机信号畸变偏差改正和工程示范验证研究，形成覆盖算法、软件、平台和应用验证的完整技术链。')
    head(d,'二、主要修改建议')
    for t in ['摘要由论文总结改为合同验收总结，突出任务完成情况和工程价值。','第二章增加现有技术不足分析，形成问题-方法-创新对应关系。','第三至第六章增加工程贡献总结，强化软件、平台和示范应用成果。','第七章改为合同指标完成情况、成果评价和推广价值总结。','全文统一术语、图表编号、标题层级和格式。']:
        add(d,t)
    head(d,'三、合同指标完成情况建议表')
    tb=d.add_table(rows=1,cols=2); tb.style='Table Grid'
    for c,t in zip(tb.rows[0].cells,['合同任务','完成情况']): c.text=t
    for a,b in [('单北斗时空基准研究','完成'),('CORS服务端改造','完成'),('电离层建模与监测','完成'),('信号畸变偏差改正','完成'),('铁路示范应用','完成'),('专利论文成果','完成')]:
        c=tb.add_row().cells; c[0].text=a; c[1].text=b
    head(d,'四、第七章优化总结')
    add(d,'课题围绕铁路单北斗差分定位关键技术开展研究，形成了单北斗时空基准、服务端处理链路、电离层质量控制、接收机偏差改正和工程验证体系，为集团公司北斗规模应用和铁路高精度位置服务提供技术支撑。')
    d.save(out/'铁路单北斗差分定位算法优化与应用研究_结题报告企业验收优化修订稿_标黄版.docx')

def make_note():
    d=Document(); d.add_paragraph('铁路单北斗差分定位算法优化与应用研究 修改说明').alignment=WD_ALIGN_PARAGRAPH.CENTER
    head(d,'一、修改原则'); add(d,'围绕企业科研课题验收要求，按照合同目标、技术成果、工程应用和推广价值重新组织报告逻辑。')
    head(d,'二、修改内容');
    for t in ['摘要重构：突出研究任务、创新成果和工程价值。','章节优化：增加问题分析、工程贡献和章节小结。','成果强化：突出算法、软件、平台和示范应用。','格式统一：规范标题、图表、术语和引用。']:
        add(d,t)
    d.save(out/'铁路单北斗差分定位算法优化与应用研究_修改说明.docx')

make_report(); make_note()
