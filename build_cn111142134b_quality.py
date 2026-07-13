from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path('output')
OUT.mkdir(exist_ok=True)

TITLE = '一种坐标时间序列处理方法及装置\n专利质量评价材料'
FILENAME = OUT / 'CN111142134B_一种坐标时间序列处理方法及装置_专利质量评价材料.docx'


def set_cn_font(run, name='宋体', size=12, bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_shading(cell, fill='D9EAF7'):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def add_paragraph(doc, text='', first_line=True, space_after=0, line_spacing=1.5, bold=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    if first_line:
        pf.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    set_cn_font(r, '宋体', 12, bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    if level == 1:
        set_cn_font(r, '黑体', 16, True)
    elif level == 2:
        set_cn_font(r, '黑体', 14, True)
    else:
        set_cn_font(r, '黑体', 12, True)
    return p


def add_numbered(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(f'{num} {text}')
    set_cn_font(r, '宋体', 12, False)
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(2.8)
sec.right_margin = Cm(2.5)

# Default styles
styles = doc.styles
styles['Normal'].font.name = '宋体'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
styles['Normal'].font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
for i, line in enumerate(TITLE.split('\n')):
    r = p.add_run(line)
    set_cn_font(r, '黑体', 18 if i == 0 else 20, True)
    if i == 0:
        r.add_break()

# Basic info table
info = doc.add_table(rows=4, cols=4)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.style = 'Table Grid'
info_data = [
    ('专利号', 'ZL201911102015.8', '授权公告号', 'CN111142134B'),
    ('专利名称', '一种坐标时间序列处理方法及装置', '授权公告日', '2022年3月11日'),
    ('专利权人', '中铁第四勘察设计院集团有限公司', '申请日', '2019年11月12日'),
    ('发明人', '马俊、曹成度、腾焕乐、闵阳、刘善勇', 'IPC主分类号', 'G01S 19/39'),
]
for row, vals in zip(info.rows, info_data):
    for j, val in enumerate(vals):
        cell = row.cells[j]
        cell.text = val
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_cn_font(run, '宋体', 10.5, j % 2 == 0)
        if j % 2 == 0:
            set_cell_shading(cell)

doc.add_paragraph()

add_heading(doc, '二、专利质量评价材料', 1)
add_heading(doc, '（一）新颖性和创造性', 2)
add_heading(doc, '1. 技术背景及所解决的技术问题', 3)
add_paragraph(doc, 'GNSS基准站坐标时间序列广泛应用于大地测量和地球动力学研究。受地球物理环境和GNSS观测系统误差等因素影响，坐标时间序列通常同时包含白噪声和有色噪声。有色噪声具有明显的时间相关性，功率主要集中于低频部分，会影响测站坐标稳定性、运动参数估值及其不确定度评定。')
add_paragraph(doc, '授权说明书指出，已有离散小波变换、小波熵、经验模态分解和奇异谱分析等方法能够削弱部分噪声，但仍存在低频有色噪声难以与周期信号或白噪声有效分离、需要预先掌握白噪声统计信息，或者难以准确估计残余有色噪声振幅等问题。针对上述不足，本专利提出一种针对同一测站北、东、天三个方向坐标时间序列的联合处理方法，以三个方向残差序列的频域功率谱为分析对象，通过标准化和主成分分析分离噪声成分，并重构新的坐标残差序列。')

add_heading(doc, '2. 最接近现有技术及对比基础', 3)
add_paragraph(doc, '本专利授权公告文本列明了CN106814378A、CN109709585A、CN105572703A、CN104392414A和CN109188466A，以及与测站坐标时间序列随机模型、噪声估计、独立成分分析、区域CORS时间序列和共模误差提取有关的论文及学位论文。上述文件表明，申请日前本领域已知的技术手段包括：利用小波变换削弱高频噪声；利用小波熵或模拟白噪声提取部分有色噪声；利用经验模态分解或奇异谱分析分离趋势、周期和噪声成分；利用主成分或独立成分分析提取多个测站之间的共模误差。')
add_paragraph(doc, '与上述技术相比，本专利并非简单地对单一坐标分量进行平滑或降噪，也不是针对多个测站提取空间共模误差，而是针对同一目标测站三个坐标分量的残差功率谱进行联合标准化、主成分划分和频域重构。该处理对象、噪声分离依据及重构流程构成了本专利新颖性和创造性评价的主要基础。')

add_heading(doc, '3. 新颖性', 3)
add_paragraph(doc, '独立权利要求1限定了一种完整的坐标时间序列处理方法。结合授权权利要求书，其主要技术特征包括：')
add_numbered(doc, '（1）', '从目标测站获取数据并进行预处理，得到北、东、天三个方向坐标时间序列对应的三个残差序列，残差由白噪声和有色噪声组成；')
add_numbered(doc, '（2）', '分别对三个残差序列进行傅里叶变换，计算三个残差序列的功率谱；')
add_numbered(doc, '（3）', '对三个方向的功率谱进行标准化处理，使原本功率谱值占比较大的方向降低对联合分析结果的主导作用；')
add_numbered(doc, '（4）', '求取三个标准残差序列功率谱的特征值和特征向量矩阵，并计算其主成分；')
add_numbered(doc, '（5）', '根据特征值将主成分划分为表征第一部分白噪声的第一主成分，以及表征有色噪声和第二部分白噪声混合成分的第二主成分；删除第二主成分并保留第一主成分；')
add_numbered(doc, '（6）', '根据更新后的主成分计算三个新的残差序列，并用新的残差序列替换原残差序列，形成新的坐标时间序列。')
add_paragraph(doc, '授权文本所列现有技术及说明书背景技术均未公开上述技术特征的完整组合。尤其是，现有技术没有同时公开：以同一测站三个方向的残差功率谱为联合分析对象；通过标准化抑制单一方向功率幅值的主导影响；依据特征值划分白噪声主成分和有色噪声混合主成分；删除后者并在频域中重构三个方向残差序列。因此，权利要求1相对于授权文本列明的现有技术具备新颖性。')
add_paragraph(doc, '权利要求2进一步限定了根据特征向量矩阵和更新后的主成分计算新的标准残差序列功率谱、进行反标准化并实施傅里叶逆变换；权利要求3和4进一步限定了相位计算以及结合原功率谱相位实施逆变换的过程。上述进一步限定同样未见于授权文本列明的现有技术。独立权利要求5以装置模组形式对应实现权利要求1的方法，权利要求6至8进一步限定反标准化、相位计算和逆变换功能，因而亦具有相应的新颖性。')

add_heading(doc, '4. 创造性', 3)
add_paragraph(doc, '以申请日前已经存在的坐标时间序列降噪方法为基础，本专利实际解决的技术问题可概括为：在三个坐标方向功率水平差异明显、低频有色噪声与有效周期信号及白噪声相互混杂的情况下，如何减弱单一方向对联合分析结果的支配作用，分离有色噪声成分，并在尽量保持测站运动参数估值稳定的同时重构适于不确定度评定的坐标时间序列。')
add_paragraph(doc, '本专利针对该问题形成了相互配合的技术方案。首先，将三个方向的残差序列变换至频域，以功率谱反映不同频率成分的能量分布；其次，对三个方向的功率谱进行标准化，避免垂直方向或其他高功率分量在主成分分析中占据绝对主导；再次，根据特征值和主成分的统计特性区分主要表征白噪声的成分与主要表征有色噪声和部分白噪声的混合成分；最后，删除混合噪声主成分，并通过特征向量反变换、反标准化、相位恢复及傅里叶逆变换重构残差序列。')
add_paragraph(doc, '上述步骤并非相互孤立。功率谱标准化为三个方向的联合主成分分析提供了可比较的数据基础；主成分划分使有色噪声成分能够从联合频谱中被识别；反标准化和相位恢复则保证处理结果能够由频域返回时间域，形成可用于替换原残差的实际坐标序列。由此形成的整体技术路线不同于单纯设置小波阈值、模拟白噪声、分解固有模态函数、提取趋势项或提取多测站共模误差的常规方法。')
add_paragraph(doc, '从授权文本所列现有技术及说明书背景技术中，不能直接得到将“同站三方向残差功率谱标准化—主成分噪声分类—删除混合噪声主成分—结合原相位重构残差序列”组合使用的明确技术启示。特别是，标准化处理解决的是三个方向量级差异导致的主成分偏置问题，主成分划分解决的是有色噪声与白噪声混合问题，相位恢复和逆变换解决的是重构结果返回时间域的问题，三者共同作用于一个完整的数据处理链。')
add_paragraph(doc, '说明书以AC62测站2011年1月至2016年12月的N、E、U坐标时间序列为实施例。经标准化和主成分处理后，重构残差功率谱的线性拟合斜率分别为0.0217、0.0243和−0.0452，均接近0；说明书据此说明新的残差序列频谱特征接近白噪声。处理前后对比还表明，测站运动参数估值变化较小，而不确定度有所减小。上述结果与本专利所要解决的技术问题相对应，能够说明区别技术特征产生了可验证的技术效果。')
add_paragraph(doc, '因此，权利要求1所述技术方案相对于授权文本列明的现有技术，不是本领域常规降噪手段的简单并列，而是围绕同站三方向坐标残差的频域联合分析和重构形成了具有内在联系的技术组合，具备突出的实质性特点和相应的技术进步。权利要求2至4进一步限定频域重构的具体过程；权利要求5至8以装置形式实现相应方法并作进一步限定，亦具备与其技术特征相适应的创造性。')

add_heading(doc, '（二）文本质量', 2)
add_heading(doc, '1. 说明书清楚、完整地公开了发明内容', 3)
add_paragraph(doc, '说明书明确记载了本专利所属的GNSS数据处理技术领域，说明了坐标时间序列中白噪声和有色噪声对测站稳定性、运动参数估计和地球物理解释的影响，并针对现有方法难以剔除低频有色噪声、难以准确估计残余噪声振幅等问题提出相应技术方案。技术问题、技术方案和技术效果之间具有明确对应关系。')
add_paragraph(doc, '说明书对主要处理步骤进行了连续、完整的公开，包括坐标时间序列预处理、运动模型建立、残差序列获取、快速傅里叶变换、功率谱计算、对数处理、功率谱标准化、方差—协方差矩阵构建、特征值和特征向量计算、主成分划分、噪声成分删除、功率谱重构、反标准化、相位计算、傅里叶逆变换和新坐标时间序列形成。相关步骤配有公式（1）至（14），各主要变量和矩阵均作了说明。')
add_paragraph(doc, '说明书附图与正文相互对应。图3概括了权利要求1的基本处理流程；图4给出了从数据预处理到运动参数及不确定度计算的完整流程；图5至图9分别展示原坐标时间序列、原残差序列、原残差功率谱、主成分处理后的功率谱以及重构后的残差序列；图10给出了装置的功能模组结构。上述内容使所属技术领域的技术人员能够理解各步骤之间的关系并据此实施。')
add_paragraph(doc, '说明书还给出了AC62测站的具体实施例，包括数据时段、粗差和跳变处理、模型参数估计、特征值结果、主成分选择依据、处理后谱斜率及处理前后参数不确定度比较。该实施例为主成分选择、功率谱重构和技术效果判断提供了具体依据，说明书公开程度能够满足实施要求。')

add_heading(doc, '2. 权利要求书清楚、简要', 3)
add_paragraph(doc, '本专利共有8项权利要求，其中权利要求1和权利要求5分别为方法独立权利要求和装置独立权利要求。权利要求1按照数据获取与预处理、频域变换、功率谱标准化、主成分分析、噪声成分删除、残差重构及坐标序列替换的顺序记载技术方案，步骤关系清楚，技术术语与说明书基本一致。')
add_paragraph(doc, '权利要求2至4分别对新的残差序列计算、频率信号相位计算以及傅里叶逆变换作进一步限定，引用关系明确。权利要求5以获取模组、计算功率谱模组、标准化模组、主成分分析模组和替换模组对应方法步骤；权利要求6至8进一步限定反标准化、相位计算和逆变换功能。方法权利要求与装置权利要求之间具有清晰的对应关系，整体层次较为合理。')

add_heading(doc, '3. 权利要求以说明书为依据，保护范围合理', 3)
add_paragraph(doc, '权利要求1中的三个方向坐标残差序列、傅里叶变换、功率谱计算与标准化、特征值和特征向量求取、主成分划分、第二主成分删除以及残差序列重构等技术特征，均可在说明书发明内容和具体实施方式中找到相应记载。说明书图3和图4对总体步骤进行了图示，公式（1）至（14）对主要计算过程进行了说明。')
add_paragraph(doc, '权利要求2所述特征向量反变换、反标准化和傅里叶逆变换，权利要求3所述频率信号相位计算，以及权利要求4所述结合相位进行逆变换，均由说明书相关段落和公式明确支持。权利要求5至8所述各功能模组亦可由方法步骤及图10所示装置结构得到。各项权利要求未明显超出说明书公开范围。')
add_paragraph(doc, '独立权利要求1围绕本专利的核心发明构思限定了从三方向残差功率谱联合分析到新坐标时间序列形成所必需的主要技术特征；从属权利要求进一步限定频域重构细节；装置权利要求以功能模组形式对应方法步骤。其保护范围与说明书公开的技术贡献总体相适应，既未仅以技术效果概括保护，也未加入与解决技术问题无关的明显非必要技术特征。')

add_heading(doc, '评价结论', 2)
add_paragraph(doc, '本专利针对GNSS测站坐标时间序列中低频有色噪声难以有效分离的问题，提出同一测站三个坐标方向残差功率谱的标准化、主成分噪声划分和频域重构方法。该技术方案相对于授权文本列明的现有技术具有明确区别，区别技术特征之间形成完整的数据处理链，并在说明书实施例中取得了与技术问题相对应的处理效果。说明书对技术问题、处理流程、计算方法和实施结果公开较为完整；权利要求层次清楚，具有说明书支持，保护范围总体合理。因此，本专利在新颖性、创造性和文本质量方面具有较好的专利质量。')

# Footnote-like source note
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.line_spacing = 1.2
r = p.add_run('材料依据：CN111142134B《一种坐标时间序列处理方法及装置》授权公告文本；本材料未引用或虚构该专利之外的论文、专利实施规模、经济效益或获奖情况。')
set_cn_font(r, '宋体', 9, False)

# Header/footer
for section in doc.sections:
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run('CN111142134B 专利质量评价材料')
    set_cn_font(hr, '宋体', 9, False)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    fp._p.append(fld)

# Keep headings with next paragraph
for p in doc.paragraphs:
    if p.text.startswith(('二、','（一）','（二）','1. ','2. ','3. ','4. ','评价结论')):
        p.paragraph_format.keep_with_next = True

doc.save(FILENAME)
print(FILENAME)
