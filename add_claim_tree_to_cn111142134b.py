from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

OUT = Path('output')
DOCX = OUT / 'CN111142134B_一种坐标时间序列处理方法及装置_专利质量评价材料.docx'
PNG = OUT / 'CN111142134B_权利要求关系图.png'


def font(size: int, bold: bool = False):
    candidates = [
        '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc' if bold else '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
        '/usr/share/fonts/opentype/noto/NotoSerifCJK-Bold.ttc' if bold else '/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc',
        '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf' if bold else '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
    ]
    for p in candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size=size)
    return ImageFont.load_default()


def centered_text(draw, box, lines, fnt, fill=(20, 20, 20), spacing=8):
    x0, y0, x1, y1 = box
    bboxes = [draw.textbbox((0, 0), s, font=fnt) for s in lines]
    heights = [b[3] - b[1] for b in bboxes]
    total_h = sum(heights) + spacing * (len(lines) - 1)
    y = y0 + (y1 - y0 - total_h) / 2
    for s, b, h in zip(lines, bboxes, heights):
        w = b[2] - b[0]
        draw.text((x0 + (x1 - x0 - w) / 2, y), s, font=fnt, fill=fill)
        y += h + spacing


def rounded(draw, box, fill, outline, width=4, radius=22):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


W, H = 1800, 1100
img = Image.new('RGB', (W, H), 'white')
d = ImageDraw.Draw(img)
line = (95, 95, 95)

# Boxes
root = (500, 70, 1300, 245)
left = (150, 390, 760, 575)
right = (1040, 390, 1650, 575)
children_y = (780, 930)
child_w = 245
xs_left = [70, 340, 610]
xs_right = [975, 1245, 1515]

rounded(d, root, (229, 242, 222), (67, 125, 52), width=4, radius=28)
rounded(d, left, (255, 239, 222), (215, 105, 25), width=4, radius=25)
rounded(d, right, (255, 239, 222), (215, 105, 25), width=4, radius=25)
for x in xs_left + xs_right:
    rounded(d, (x, children_y[0], x + child_w, children_y[1]), (229, 239, 252), (46, 99, 184), width=4, radius=22)

# Connectors
root_cx = (root[0] + root[2]) // 2
split_y = 330
left_cx = (left[0] + left[2]) // 2
right_cx = (right[0] + right[2]) // 2
d.line((root_cx, root[3], root_cx, split_y), fill=line, width=5)
d.line((left_cx, split_y, right_cx, split_y), fill=line, width=5)
d.line((left_cx, split_y, left_cx, left[1]), fill=line, width=5)
d.line((right_cx, split_y, right_cx, right[1]), fill=line, width=5)

branch_y = 690
for parent_cx, xs in [(left_cx, xs_left), (right_cx, xs_right)]:
    d.line((parent_cx, left[3], parent_cx, branch_y), fill=line, width=5)
    first_cx = xs[0] + child_w // 2
    last_cx = xs[-1] + child_w // 2
    d.line((first_cx, branch_y, last_cx, branch_y), fill=line, width=5)
    for x in xs:
        cx = x + child_w // 2
        d.line((cx, branch_y, cx, children_y[0]), fill=line, width=5)

# Text
centered_text(d, root, ['一种坐标时间序列处理方法及装置'], font(52, True))
centered_text(d, left, ['独立权利要求1', '坐标时间序列处理方法'], font(38, True))
centered_text(d, right, ['独立权利要求5', '坐标时间序列处理装置'], font(38, True))
for x, n in zip(xs_left, [2, 3, 4]):
    centered_text(d, (x, children_y[0], x + child_w, children_y[1]), [f'权利要求{n}'], font(34, True))
for x, n in zip(xs_right, [6, 7, 8]):
    centered_text(d, (x, children_y[0], x + child_w, children_y[1]), [f'权利要求{n}'], font(34, True))

img.save(PNG, quality=95)

doc = Document(DOCX)

# Find the paragraph introducing the next subsection, and insert the figure before it.
target = None
for p in doc.paragraphs:
    if p.text.strip().startswith('3. 权利要求以说明书为依据'):
        target = p
        break

heading = doc.add_paragraph()
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
heading.paragraph_format.space_before = Pt(8)
heading.paragraph_format.space_after = Pt(5)
r = heading.add_run('本专利权利要求关系如下图所示。')
r.font.name = '宋体'
r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
r.font.size = Pt(12)

pic = doc.add_paragraph()
pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic.paragraph_format.space_after = Pt(3)
pic.add_run().add_picture(str(PNG), width=Cm(15.6))

cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_after = Pt(8)
r = cap.add_run('图2-1  本专利权利要求关系图')
r.font.name = '宋体'
r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
r.font.size = Pt(10.5)

if target is not None:
    target._p.addprevious(heading._p)
    target._p.addprevious(pic._p)
    target._p.addprevious(cap._p)

# If no target was found, the paragraphs remain appended at the end.
doc.save(DOCX)
print(DOCX)
